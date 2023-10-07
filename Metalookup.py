import argparse
import os
from PyPDF2 import PdfReader
from PIL import Image, UnidentifiedImageError, ExifTags
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from hachoir.metadata import extractMetadata
from hachoir.parser import createParser

Version = "0.2.0"

Banner = r"""
  __  __     _          _            _             
 |  \/  |___| |_ __ _  | |   ___  __| |___  _ _ __ 
 | |\/| / -_)  _/ _` | | |__/ _ \/ _| / / || | '_ \
 |_|  |_\___|\__\__,_| |____\___/\__|_\_\\_,_| .__/
                                             |_|   
"""


def safe_getattr(obj, attr, default=None):
    return getattr(obj, attr, default)

def extract_pdf_metadata(pdf_path):
    with open(pdf_path, 'rb') as file:
        pdf = PdfReader(file)
        info = pdf.trailer["/Info"]
        metadata = {key[1:]: info[key] for key in info if key != '/ID'}
        return metadata

def extract_video_metadata(video_path):
    parser = createParser(video_path)
    if not parser:
        print(f"Unable to parse video file: {video_path}")
        return {}
    
    with parser:
        try:
            metadata = extractMetadata(parser)
        except Exception as e:
            print(f"Metadata extraction error: {e}")
            return {}
    
    if not metadata:
        return {}

    # Iterate over metadata attributes and fetch their values
    meta_data_dict = {}
    for item in metadata:
        if item.values:
            meta_data_dict[item.key] = item.values[0].value

    return meta_data_dict


def extract_image_metadata(img_path):
    try:
        with Image.open(img_path) as img:
            info = img.info
            
            # Check if the image has the "exif" attribute
            if "exif" in info and info["exif"]:
                exif_data = img._getexif()
                
                if exif_data:
                    # Convert the tag ID to the tag name for better readability
                    decoded_exif = {ExifTags.TAGS[k]: v for k, v in exif_data.items() if k in ExifTags.TAGS}
                    info["decoded_exif"] = decoded_exif

            # Remove raw EXIF data from the metadata
            info.pop("exif", None)
            
            metadata = {key: info[key] for key in info if not key.startswith('icc_')}
            return metadata
            
    except UnidentifiedImageError:
        print(f"Error: Could not identify image: {img_path}")
        return {}

def extract_office_metadata(doc_path):
    if doc_path.endswith('.docx'):
        doc = Document(doc_path)
        core_props = doc.core_properties
    elif doc_path.endswith('.xlsx'):
        wb = load_workbook(doc_path)
        core_props = wb.properties
    elif doc_path.endswith('.pptx'):
        pres = Presentation(doc_path)
        core_props = pres.core_properties
    else:
        return {}

    metadata = {
        'author': safe_getattr(core_props, 'author'),
        'category': safe_getattr(core_props, 'category'),
        'comments': safe_getattr(core_props, 'comments'),
        'content_status': safe_getattr(core_props, 'content_status'),
        'created': safe_getattr(core_props, 'created'),
        'identifier': safe_getattr(core_props, 'identifier'),
        'keywords': safe_getattr(core_props, 'keywords'),
        'language': safe_getattr(core_props, 'language'),
        'last_modified_by': safe_getattr(core_props, 'last_modified_by'),
        'last_printed': safe_getattr(core_props, 'last_printed'),
        'modified': safe_getattr(core_props, 'modified'),
        'revision': safe_getattr(core_props, 'revision'),
        'subject': safe_getattr(core_props, 'subject'),
        'title': safe_getattr(core_props, 'title'),
        'version': safe_getattr(core_props, 'version')
    }
    return metadata

def extract_metadata(file_path):
    if file_path.endswith('.pdf'):
        return extract_pdf_metadata(file_path)
    elif file_path.lower().endswith(('.png', '.jpg', '.jpeg', '.tiff', '.bmp', '.gif')):
        return extract_image_metadata(file_path)
    elif file_path.lower().endswith(('.docx', '.xlsx', '.pptx')):
        return extract_office_metadata(file_path)
    elif file_path.lower().endswith(('.mp4', '.mkv')):
        return extract_video_metadata(file_path)
    else:
        print(f"Unsupported file format for {file_path}")
        return {}

def extract_metadata_from_directory(directory):
    for root, _, files in os.walk(directory):
        for name in files:
            if name.startswith('~$'):
                continue  # Skip files starting with ~$
            file_path = os.path.join(root, name)
            metadata = extract_metadata(file_path)
            print(f"Metadata for {file_path}:\n{metadata}\n")

if __name__ == '__main__':
    print(Banner)
    parser = argparse.ArgumentParser(description='Extract metadata from files.')
    parser.add_argument('-v', '--version', action='version', version=f"Metalookup Version {Version}")
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument('-f', '--file', help='File to extract metadata from', type=str)
    group.add_argument('-d', '--directory', help='Directory to extract metadata from all contained files', type=str)

    args = parser.parse_args()

    if args.file:
        if not os.path.exists(args.file):
            print(f"Error: The file '{args.file}' does not exist.")
        else:
            metadata = extract_metadata(args.file)
            print(f"Metadata for {args.file}:\n{metadata}\n")
    elif args.directory:
        extract_metadata_from_directory(args.directory)

